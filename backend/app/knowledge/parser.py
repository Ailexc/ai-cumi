"""Document parser: extract text/markdown from local files."""
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Dict, Any


@dataclass
class ParsedDocument:
    content: str
    metadata: Dict[str, Any] = field(default_factory=dict)
    page_contents: List[str] = field(default_factory=list)


class DocumentParser:
    """Unified document parser."""

    def parse(self, file_path: str, mime_type: str | None = None) -> ParsedDocument:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = path.suffix.lower()
        if ext in {".md", ".txt", ".py", ".js", ".ts", ".tsx", ".jsx", ".json", ".yaml", ".yml", ".html", ".css", ".sql", ".sh", ".bash"}:
            return self._parse_text(path)
        if ext == ".pdf":
            return self._parse_pdf(path)
        if ext in {".docx", ".doc"}:
            return self._parse_docx(path)
        if ext in {".html", ".htm"}:
            return self._parse_html(path)
        raise ValueError(f"Unsupported file type: {ext}")

    def _parse_text(self, path: Path) -> ParsedDocument:
        content = path.read_text(encoding="utf-8", errors="replace")
        return ParsedDocument(content=content, metadata={"source": path.name})

    def _parse_pdf(self, path: Path) -> ParsedDocument:
        try:
            import fitz
        except ImportError:
            raise RuntimeError("pymupdf not installed")

        doc = fitz.open(str(path))
        pages: List[str] = []
        try:
            for page in doc:
                text = page.get_text()
                pages.append(text)
        finally:
            doc.close()

        full = "\n\n".join(pages)
        return ParsedDocument(
            content=full,
            metadata={"source": path.name, "page_count": len(pages)},
            page_contents=pages,
        )

    def _parse_docx(self, path: Path) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError:
            raise RuntimeError("python-docx not installed")

        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full = "\n\n".join(paragraphs)
        return ParsedDocument(
            content=full,
            metadata={"source": path.name, "paragraph_count": len(paragraphs)},
        )

    def _parse_html(self, path: Path) -> ParsedDocument:
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            raise RuntimeError("beautifulsoup4 not installed")

        html = path.read_text(encoding="utf-8", errors="replace")
        soup = BeautifulSoup(html, "html.parser")
        # Remove script/style
        for tag in soup(["script", "style"]):
            tag.decompose()
        text = soup.get_text(separator="\n")
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        return ParsedDocument(content="\n".join(lines), metadata={"source": path.name})
